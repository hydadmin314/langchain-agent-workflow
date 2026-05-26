from __future__ import annotations

import hashlib
import re
import zipfile
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

import fitz
from docx import Document


SUPPORTED_FILE_TYPES = {"txt", "docx", "xlsx", "pdf"}
PENDING_FILE_TYPES = {"doc", "xls", "jpg", "jpeg", "png"}


@dataclass(frozen=True)
class DocumentBlock:
    block_id: str
    block_type: str
    text: str
    source_location: dict[str, Any]
    metadata: dict[str, Any]

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


@dataclass(frozen=True)
class LoadedDocument:
    document_id: str
    source_path: str
    filename: str
    source_file_type: str
    source_file_hash: str
    modified_at: str
    blocks: list[DocumentBlock]
    warnings: list[str]

    def to_context(self, *, max_chars: int = 60000) -> str:
        parts: list[str] = []
        total = 0
        for block in self.blocks:
            item = {
                "block_id": block.block_id,
                "block_type": block.block_type,
                "source_location": block.source_location,
                "text": block.text,
            }
            rendered = repr(item)
            if total + len(rendered) > max_chars:
                parts.append(
                    repr(
                        {
                            "block_id": "context_truncated",
                            "block_type": "warning",
                            "source_location": {},
                            "text": "文档内容较长，后续块已被截断。需要完整抽取时请提高 max_context_chars 或分批抽取。",
                        }
                    )
                )
                break
            parts.append(rendered)
            total += len(rendered)
        return "\n".join(parts)


class DocumentLoader:
    """Load business documents into location-aware text blocks for LLM extraction."""

    def load(self, file_path: str | Path) -> LoadedDocument:
        path = Path(file_path).resolve()
        file_type = path.suffix.lower().lstrip(".")
        file_hash = sha256_file(path)
        stat = path.stat()
        warnings: list[str] = []

        if file_type == "txt":
            blocks = self._load_txt(path)
        elif file_type == "docx":
            blocks = self._load_docx(path)
        elif file_type == "xlsx":
            blocks = self._load_xlsx(path)
        elif file_type == "pdf":
            blocks = self._load_pdf(path)
        elif file_type in PENDING_FILE_TYPES:
            warnings.append(f"{file_type} 深度解析暂未启用，后续可接 OCR 或格式转换。")
            blocks = []
        else:
            warnings.append(f"不支持的文件类型：{file_type}")
            blocks = []

        document_id = build_document_id(path, file_hash)
        return LoadedDocument(
            document_id=document_id,
            source_path=str(path),
            filename=path.name,
            source_file_type=file_type,
            source_file_hash=file_hash,
            modified_at=datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
            blocks=blocks,
            warnings=warnings,
        )

    def _load_txt(self, path: Path) -> list[DocumentBlock]:
        raw = path.read_bytes()
        text = decode_text(raw)
        lines = [clean(line) for line in text.splitlines() if clean(line)]
        return [
            DocumentBlock(
                block_id=f"txt_line_{index:04d}",
                block_type="text_line",
                text=line,
                source_location={"path": str(path), "paragraph_index": index},
                metadata={},
            )
            for index, line in enumerate(lines)
        ]

    def _load_docx(self, path: Path) -> list[DocumentBlock]:
        document = Document(str(path))
        blocks: list[DocumentBlock] = []

        for index, paragraph in enumerate(document.paragraphs):
            text = clean(paragraph.text)
            if not text:
                continue
            blocks.append(
                DocumentBlock(
                    block_id=f"p_{index:04d}",
                    block_type="paragraph",
                    text=text,
                    source_location={"path": str(path), "paragraph_index": index},
                    metadata={},
                )
            )

        for table_index, table in enumerate(document.tables):
            for row_index, row in enumerate(table.rows):
                cells = [clean(cell.text) for cell in row.cells]
                cells = collapse_repeated_cells(cells)
                if not any(cells):
                    continue
                text = " | ".join(cells)
                blocks.append(
                    DocumentBlock(
                        block_id=f"t_{table_index:03d}_r_{row_index:04d}",
                        block_type="table_row",
                        text=text,
                        source_location={
                            "path": str(path),
                            "table_index": table_index,
                            "row_index": row_index,
                        },
                        metadata={"cells": cells},
                    )
                )
        return blocks

    def _load_xlsx(self, path: Path) -> list[DocumentBlock]:
        blocks: list[DocumentBlock] = []
        with zipfile.ZipFile(path) as zf:
            shared_strings = read_shared_strings(zf)
            for sheet_path in sorted(name for name in zf.namelist() if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")):
                sheet_name = Path(sheet_path).stem
                for row in read_sheet_rows(zf, sheet_path, shared_strings):
                    text = " | ".join(row["cells"])
                    blocks.append(
                        DocumentBlock(
                            block_id=f"{sheet_name}_r_{row['row_index']:04d}",
                            block_type="sheet_row",
                            text=text,
                            source_location={
                                "path": str(path),
                                "sheet_name": sheet_name,
                                "row_index": row["row_index"],
                            },
                            metadata={"cells": row["cells"]},
                        )
                    )
        return blocks

    def _load_pdf(self, path: Path) -> list[DocumentBlock]:
        blocks: list[DocumentBlock] = []
        with fitz.open(path) as document:
            for page_index, page in enumerate(document):
                text = clean(page.get_text("text"))
                if not text:
                    continue
                blocks.append(
                    DocumentBlock(
                        block_id=f"pdf_page_{page_index + 1:04d}",
                        block_type="pdf_page",
                        text=text,
                        source_location={"path": str(path), "page": page_index + 1},
                        metadata={},
                    )
                )
        return blocks


def build_document_id(path: Path, file_hash: str) -> str:
    key = f"{path.resolve()}|{file_hash}".encode("utf-8", errors="ignore")
    return "doc_" + hashlib.sha1(key).hexdigest()[:12]


def sha256_file(path: Path, chunk_size: int = 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        while True:
            chunk = file.read(chunk_size)
            if not chunk:
                break
            digest.update(chunk)
    return digest.hexdigest()


def decode_text(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk", "big5"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def clean(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\u3000", " ")).strip()


def collapse_repeated_cells(cells: list[str]) -> list[str]:
    result: list[str] = []
    previous: str | None = None
    for cell in cells:
        if cell == previous:
            continue
        result.append(cell)
        previous = cell
    return result


def read_shared_strings(zf: zipfile.ZipFile) -> list[str]:
    if "xl/sharedStrings.xml" not in zf.namelist():
        return []
    root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    values: list[str] = []
    for item in root.findall(f".//{ns}si"):
        texts = [node.text or "" for node in item.findall(f".//{ns}t")]
        values.append("".join(texts))
    return values


def read_sheet_rows(zf: zipfile.ZipFile, sheet_path: str, shared_strings: list[str]) -> list[dict[str, Any]]:
    root = ET.fromstring(zf.read(sheet_path))
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    rows: list[dict[str, Any]] = []
    for row in root.findall(f".//{ns}sheetData/{ns}row"):
        cells: list[str] = []
        for cell in row.findall(f"{ns}c"):
            cells.append(clean(cell_value(cell, shared_strings, ns)))
        while cells and not cells[-1]:
            cells.pop()
        if any(cells):
            rows.append({"row_index": int(row.attrib.get("r", "0")), "cells": cells})
    return rows


def cell_value(cell: ET.Element, shared_strings: list[str], ns: str) -> str:
    cell_type = cell.attrib.get("t")
    value = cell.find(f"{ns}v")
    inline = cell.find(f"{ns}is/{ns}t")
    if inline is not None:
        return inline.text or ""
    if value is None:
        return ""
    raw = value.text or ""
    if cell_type == "s":
        try:
            return shared_strings[int(raw)]
        except (ValueError, IndexError):
            return raw
    return raw
