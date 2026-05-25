import json
import re
from datetime import datetime
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[2]
DATA_DIR = ROOT / "data" / "raw"
OUTPUT_DIR = ROOT / "data" / "product_doc_agent" / "output"

DOCUMENT_TITLE_KEYWORDS = (
    "申请登记表",
    "登记表",
    "业务申请表",
    "业务受理单",
    "服务协议",
    "产品说明",
    "营销活动",
    "套餐",
)
APPLICATION_SECTION_KEYWORDS = (
    "客户信息",
    "用户信息",
    "企业信息",
    "办理信息",
    "申请信息",
    "受理信息",
    "基础套餐申请信息",
    "业务申请信息",
)
BUSINESS_SECTION_KEYWORDS = (
    "基础套餐",
    "套餐信息",
    "业务信息",
    "业务申请信息",
    "申请信息",
    "产品信息",
    "资费",
    "费用",
    "权益",
    "增值",
)
APPLICATION_FIELD_KEYWORDS = (
    "客户",
    "企业",
    "单位",
    "联系人",
    "联系电话",
    "手机",
    "地址",
    "证件",
    "身份证",
    "代码",
    "邮编",
    "行业",
    "规模",
    "经办人",
    "邮箱",
    "E-MAIL",
    "Email",
    "传真",
)
APPLICATION_VALUE_WORDS = (
    "同公司地址",
    "其它",
    "其他",
    "托收",
    "托 收",
    "现金",
    "现 金",
)
PROVIDER_PATTERNS = (
    r"中国电信[^，。；\s]*",
    r"中国联通[^，。；\s]*",
    r"中国联合网络通信[^，。；\s]*",
    r"中国移动[^，。；\s]*",
    r"[^，。；\s]*(?:电信|联通|移动|通信|通讯|信息|网络)[^，。；\s]*(?:公司|分公司|营业厅)?",
)
FEE_KEYWORDS = ("费用", "资费", "收费", "价格", "金额", "一次性", "新装", "调试费", "安装费", "手续费", "月租费")
OPTIONAL_KEYWORDS = ("可选", "增值", "权益", "赠送", "免费", "免收", "配套", "附加", "升级包", "加装")
BASE_PACKAGE_KEYWORDS = ("基础套餐", "主套餐", "套餐", "产品资费", "标准资费")
BASE_PACKAGE_ROW_KEYWORDS = ("基础套餐", "主套餐", "产品资费", "标准资费")
INCLUDED_ITEM_KEYWORDS = ("套餐内", "包含内容", "含内容", "包含服务", "内含", "随套餐")
AGREEMENT_RULE_KEYWORDS = ("营销规则", "客户特别关注", "协议", "违约", "限制", "退订", "注销", "拆机", "售后", "服务条款")
SUPPORTING_FORM_KEYWORDS = ("报备表", "备案", "承诺书", "责任人", "附录", "附件", "信息安全")
SPEED_PATTERN = r"\d+(?:\.\d+)?\s*(?:M|G|Mbps|Gbps|兆)(?:\s*/\s*\d+(?:\.\d+)?\s*(?:M|G|Mbps|Gbps|兆))?"
PRICE_PATTERN = r"\d+(?:\.\d+)?\s*元\s*/\s*(?:月|年|2年|半年|季度)(?:\s*/\s*线)?"


def clean(text):
    return re.sub(r"\s+", " ", (text or "").replace("\u3000", " ")).strip()


def dedupe(items):
    result = []
    seen = set()
    for item in items:
        key = json.dumps(item, ensure_ascii=False, sort_keys=True) if isinstance(item, dict) else item
        if key not in seen:
            seen.add(key)
            result.append(item)
    return result


def row_cells(row):
    cells = []
    previous = None
    for cell in row.cells:
        text = clean(cell.text)
        if text != previous:
            cells.append(text)
            previous = text
    return cells


def get_rows(table):
    return [row_cells(row) for row in table.rows]


def paragraphs(doc):
    return [clean(p.text) for p in doc.paragraphs if clean(p.text)]


def first_match(pattern, text, default=""):
    match = re.search(pattern, text)
    return match.group(1).strip() if match else default


def first_non_empty(values):
    return next((value for value in values if clean(value)), "")


def looks_like_issuer(text):
    return bool(text and any(re.search(pattern, text) for pattern in PROVIDER_PATTERNS))


def extract_issuer(paras):
    for para in paras[:8]:
        if looks_like_issuer(para):
            return para
    return ""


def extract_title(path, paras):
    candidates = [para for para in paras if any(keyword in para for keyword in DOCUMENT_TITLE_KEYWORDS)]
    candidates.extend(para for para in paras if para and not looks_like_issuer(para))
    return first_non_empty(candidates) or path.stem


def strip_provider_name(text, issuer=""):
    cleaned = text
    if issuer:
        cleaned = cleaned.replace(issuer, "")
    for pattern in PROVIDER_PATTERNS:
        cleaned = re.sub(pattern, "", cleaned)
    return clean(cleaned)


def extract_effective_from(path, title):
    text = f"{path.name} {title}"
    return (
        first_match(r"【(.+?)起】", text)
        or first_match(r"(\d{4}[./-]\d{1,2}[./-]\d{1,2})\s*起", text)
        or first_match(r"(\d{4}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日)\s*起", text)
        or first_match(r"(\d{4}[./-]\d{1,2}[./-]\d{1,2})", text)
    )


def extract_version(title):
    return (
        first_match(r"[（(]([^（）()]*\d{4}[^（）()]*)[）)]", title)
        or first_match(r"(?:版本|版次|编号)[:：]?\s*([A-Za-z0-9_.\-/]+)", title)
    )


def extract_meta(path, paras):
    path = Path(path).resolve()
    try:
        source_path = str(path.relative_to(ROOT))
    except ValueError:
        source_path = str(path)
    title = extract_title(path, paras)
    issuer = extract_issuer(paras)
    version = extract_version(title)
    effective_from = extract_effective_from(path, title)
    product_name = re.split(r"申请登记表|登记表|业务申请表|业务受理单|服务协议|产品说明", title)[0]
    product_name = strip_provider_name(product_name, issuer)
    return {
        "filename": path.name,
        "source_path": source_path,
        "issuer": issuer,
        "title": title,
        "version": version,
        "effective_from": effective_from,
        "product_name": product_name or title,
    }


def normalize_field_label(text):
    label = clean(text).replace("*", "").replace(" ", "")
    label = re.sub(r"[:：]+$", "", label)
    return label


def is_application_value_cell(text):
    normalized = clean(text).replace(" ", "")
    if not normalized:
        return False
    return normalized in {value.replace(" ", "") for value in APPLICATION_VALUE_WORDS}


def is_application_field_cell(row, index):
    cell = clean(row[index]) if index < len(row) else ""
    label = normalize_field_label(cell)
    if not label or len(label) > 24:
        return False
    if cell.startswith("□") or "银行账号" in label or is_application_value_cell(cell):
        return False
    if "*" in cell:
        return True
    previous = clean(row[index - 1]) if index > 0 else ""
    next_cell = clean(row[index + 1]) if index + 1 < len(row) else ""
    has_field_keyword = any(keyword in label for keyword in APPLICATION_FIELD_KEYWORDS)
    if not has_field_keyword:
        return False
    return index == 0 or previous == "" or next_cell == ""


def split_plain_options(text):
    normalized = clean(text)
    if not normalized or normalized.startswith("□"):
        return []
    option_words = []
    for word in APPLICATION_VALUE_WORDS:
        pattern = re.escape(word)
        if re.search(pattern, normalized):
            option_words.append(clean(word))
    return dedupe(option_words)


def extract_application_fields(rows):
    fields = []
    for row_index, row in enumerate(rows):
        if row and any(keyword in " ".join(row) for keyword in BUSINESS_SECTION_KEYWORDS):
            break
        field_indexes = [
            index
            for index, _ in enumerate(row)
            if is_application_field_cell(row, index)
        ]
        for position, index in enumerate(field_indexes):
            cell = row[index]
            label = normalize_field_label(cell)
            if not label:
                continue
            next_index = field_indexes[position + 1] if position + 1 < len(field_indexes) else len(row)
            value_cells = [clean(value) for value in row[index + 1:next_index] if clean(value)]
            options = []
            for value in value_cells:
                options.extend(checkbox_options(value))
                options.extend(split_plain_options(value))
            plain_values = [
                value
                for value in value_cells
                if not value.startswith("□") and not checkbox_options(value) and not split_plain_options(value)
            ]
            fields.append({
                "label": label,
                "required": "*" in cell,
                "value": " ".join(plain_values),
                "options": dedupe(options),
                "source_evidence": " | ".join(row),
                "row_index": row_index,
            })
    return dedupe(fields)


def extract_required_fields(rows):
    return extract_application_fields(rows)


def checkbox_options(text):
    if "□" not in (text or ""):
        return []
    options = []
    for part in re.split(r"□", text):
        part = clean(part)
        if not part:
            continue
        part = re.split(r"(?=□)|。|；", part)[0].strip(" ，,")
        if part:
            options.append(part)
    return dedupe(options)


def parse_base_packages(text):
    packages = []
    speed_pattern = f"({SPEED_PATTERN})"
    price_pattern = r"(\d+(?:\.\d+)?)\s*元\s*/\s*(月|年|2年|半年|季度)(?:\s*/\s*线)?"
    for speed, amount, period in re.findall(speed_pattern + r".{0,30}?" + price_pattern, text, flags=re.IGNORECASE):
        numeric_amount = float(amount)
        packages.append({
            "speed": clean(speed),
            "amount": int(numeric_amount) if numeric_amount.is_integer() else numeric_amount,
            "currency": "CNY",
            "billing_period": period,
            "raw": f"{clean(speed)} {amount}元/{period}",
        })
    return dedupe(packages)


def parse_package_prices_for_speed(text, speed, package_name=""):
    packages = []
    if not speed:
        return packages
    start = text.find(speed)
    segment = text[start:] if start >= 0 else text
    next_speed = re.search(SPEED_PATTERN, segment[len(speed):], flags=re.IGNORECASE)
    if next_speed:
        segment = segment[: len(speed) + next_speed.start()]
    for amount, period in re.findall(r"(\d+(?:\.\d+)?)\s*元\s*/\s*(月|年|2年|半年|季度)(?:\s*/\s*线)?", segment):
        numeric_amount = float(amount)
        packages.append({
            "name": package_name,
            "speed": clean(speed),
            "amount": int(numeric_amount) if numeric_amount.is_integer() else numeric_amount,
            "currency": "CNY",
            "billing_period": period,
            "raw": f"{clean(speed)} {amount}元/{period}",
        })
    return packages


def parse_base_package_rows(rows):
    packages = []
    for item in rows:
        text = clean(" ".join(str(item.get(key, "")) for key in ("name", "description")))
        paired_packages = parse_base_packages(text)
        if paired_packages:
            for package in paired_packages:
                package.setdefault("name", item.get("name", ""))
            packages.extend(paired_packages)
            continue
        speeds = [clean(match.group(0)) for match in re.finditer(SPEED_PATTERN, text, flags=re.IGNORECASE)]
        if speeds:
            row_packages = []
            for speed in dedupe(speeds):
                row_packages.extend(parse_package_prices_for_speed(text, speed, item.get("name", "")))
            if row_packages:
                packages.extend(row_packages)
                continue
        packages.extend(parse_base_packages(text))
    return dedupe(packages)


def parse_variable_package_rows(business_rows):
    packages = []
    for item in business_rows:
        name = item["name"] or item["category"]
        description = item["description"]
        if not is_customer_input_package_row(item):
            continue
        speed_range = first_match(r"[（(]([^（）()]*M[^（）()]*)[）)]", description)
        packages.append({
            "name": name,
            "speed": speed_range or "需填写",
            "amount": None,
            "currency": "CNY",
            "billing_period": "月",
            "raw": description,
            "pricing_type": "customer_input",
        })
    return packages


def parse_uplink_packages(text):
    packages = []
    current_speed = ""
    for token in re.split(r"\s+", text):
        speed = re.search(r"上行升速(\d+M)", token)
        if speed:
            current_speed = speed.group(1)
            continue
        price = re.search(r"(\d+)元/(月|年)", token)
        if current_speed and price:
            packages.append({
                "name": f"上行升速{current_speed}",
                "amount": int(price.group(1)),
                "currency": "CNY",
                "billing_period": price.group(2),
            })
    return packages


def is_section_header(row):
    joined = " ".join(row)
    if any(keyword in joined for keyword in APPLICATION_SECTION_KEYWORDS + BUSINESS_SECTION_KEYWORDS):
        return True
    return len(row) <= 2 and bool(joined) and len(joined) <= 30


def has_business_payload(row):
    joined = " ".join(row)
    return (
        len(row) >= 3
        or "□" in joined
        or re.search(r"\d+(?:\.\d+)?\s*元", joined)
        or any(keyword in joined for keyword in FEE_KEYWORDS + OPTIONAL_KEYWORDS + INCLUDED_ITEM_KEYWORDS)
    )


def classify_business_rows(rows):
    application_rows = []
    in_business = not any(any(keyword in " ".join(row) for keyword in BUSINESS_SECTION_KEYWORDS) for row in rows)
    for index, row in enumerate(rows):
        joined = " ".join(row)
        if any(keyword in joined for keyword in BUSINESS_SECTION_KEYWORDS):
            in_business = True
            if not has_business_payload(row):
                continue
        if not in_business or not row:
            continue
        if any(keyword in joined for keyword in APPLICATION_SECTION_KEYWORDS) and not any(keyword in joined for keyword in BUSINESS_SECTION_KEYWORDS):
            continue
        category = clean(row[0])
        name = clean(row[1]) if len(row) > 2 else category
        description = clean(" ".join(row[2:] if len(row) > 2 else row[1:]))
        if not description and len(row) == 2:
            description = clean(row[1])
        if not description and len(row) == 1:
            description = category
        if is_section_header(row) and (len(row) == 1 or (len(row) <= 2 and not has_business_payload(row))):
            continue
        if category == "填表说明":
            category = "说明"
        application_rows.append({
            "row_index": index,
            "category": category,
            "name": name,
            "description": description,
            "options": checkbox_options(description),
        })
    return application_rows


def extract_sla(tables, business_rows):
    sla = {"eligibility": "", "plans": [], "compensation_rules": []}
    for item in business_rows:
        if "SLA" in item["name"] or "SLA" in item["category"]:
            sla["eligibility"] = item["description"]
    for table in tables:
        rows = get_rows(table)
        if not rows:
            continue
        if rows[0] and rows[0][0] == "产品" and len(rows) >= 2:
            headers = rows[0][1:] or ["标准"]
            for plan_index, header in enumerate(headers):
                plan = {"product": header, "metrics": {}}
                for row in rows[1:]:
                    if len(row) > plan_index + 1:
                        plan["metrics"][row[0]] = row[plan_index + 1]
                sla["plans"].append(plan)
        if rows[0] and "月可用率范围" in rows[0][0]:
            for row in rows[1:]:
                if len(row) >= 2:
                    sla["compensation_rules"].append({
                        "availability_range": row[0],
                        "compensation": row[1],
                    })
    return sla


def extract_supporting_forms(tables):
    forms = []
    for idx, table in enumerate(tables):
        rows = get_rows(table)
        if not rows:
            continue
        title = clean(" ".join(rows[0]))
        if "IP地址报备表" in title:
            forms.append({"table_index": idx, "name": "IP地址报备表", "required_fields": extract_required_fields(rows)})
        elif "ICP备案号" in title:
            forms.append({"table_index": idx, "name": "域名备案信息表", "columns": rows[0]})
        elif "APP下载URL" in title:
            forms.append({"table_index": idx, "name": "APP应用备案信息表", "columns": rows[0]})
        elif any(keyword in title for keyword in SUPPORTING_FORM_KEYWORDS):
            forms.append({
                "table_index": idx,
                "name": title,
                "required_fields": extract_required_fields(rows),
                "columns": rows[0],
            })
    return forms


def extract_marketing_rules(paras):
    start = 0
    for index, para in enumerate(paras):
        if any(keyword in para for keyword in AGREEMENT_RULE_KEYWORDS):
            start = index + 1
            break
    rules = []
    for para in paras[start:]:
        if any(keyword in para for keyword in DOCUMENT_TITLE_KEYWORDS) or len(para) < 12:
            continue
        rules.append(para)
    return rules


def flatten_table_rows(tables):
    rows = []
    for table in tables:
        rows.extend(get_rows(table))
    return rows


def row_matches(item, keywords):
    text = " ".join(str(item.get(key, "")) for key in ("category", "name", "description"))
    return any(keyword in text for keyword in keywords)


def row_header_matches(item, keywords):
    text = " ".join(str(item.get(key, "")) for key in ("category", "name"))
    return any(keyword in text for keyword in keywords)


def row_text(item):
    return clean(" ".join(str(item.get(key, "")) for key in ("category", "name", "description")))


def has_price(text):
    return bool(re.search(PRICE_PATTERN, text))


def has_placeholder_price(text):
    return bool(re.search(r"(月租费|费用|资费|收费)\s*[：:]\s*[_＿]{2,}", text))


def is_base_package_row(item):
    header = clean(" ".join(str(item.get(key, "")) for key in ("category", "name")))
    text = row_text(item)
    if row_header_matches(item, INCLUDED_ITEM_KEYWORDS) or row_header_matches(item, OPTIONAL_KEYWORDS):
        return False
    if any(keyword in header for keyword in BASE_PACKAGE_ROW_KEYWORDS) and (has_price(text) or has_placeholder_price(text)):
        return True
    if has_placeholder_price(text) and re.search(SPEED_PATTERN, text, flags=re.IGNORECASE):
        return True
    return False


def is_customer_input_package_row(item):
    text = row_text(item)
    if row_header_matches(item, OPTIONAL_KEYWORDS) or row_header_matches(item, INCLUDED_ITEM_KEYWORDS):
        return False
    return has_placeholder_price(text) and re.search(SPEED_PATTERN, text, flags=re.IGNORECASE)


def is_optional_service_row(item):
    if is_base_package_row(item):
        return False
    header_matches = row_header_matches(item, OPTIONAL_KEYWORDS)
    text = row_text(item)
    if header_matches:
        return True
    if row_header_matches(item, INCLUDED_ITEM_KEYWORDS):
        return False
    return any(keyword in text for keyword in ("上行升速包", "IP地址升级", "付费升级", "可升级", "增值服务", "月基本费0元", "申请线数", "配套业务"))


def is_fee_rule_row(item):
    if is_base_package_row(item):
        return False
    text = row_text(item)
    header = clean(" ".join(str(item.get(key, "")) for key in ("category", "name")))
    if row_header_matches(item, INCLUDED_ITEM_KEYWORDS):
        return False
    if any(keyword in header for keyword in ("一次性", "新装优惠", "安装", "调测", "调试", "手续费", "费用", "期限", "协议期")):
        return True
    if any(keyword in text for keyword in ("一次性费用", "安装调测费", "手续费", "协议期", "违约金", "按天折算", "未注明的资费")):
        return True
    return False


def is_blank_business_row(item):
    return not clean(row_text(item)).strip("| ")


def is_supporting_form_business_row(item):
    return any(keyword in row_text(item) for keyword in SUPPORTING_FORM_KEYWORDS)


def service_attribute_type(item):
    text = row_text(item)
    header = clean(" ".join(str(item.get(key, "")) for key in ("category", "name")))
    if "接口" in header or "接口" in text:
        return "interface_standard"
    if "套餐类型" in header:
        return "package_type"
    if "SLA" in text or "服务等级" in text:
        return "sla_service"
    if "专享" in header or "专享" in text:
        return "exclusive_service"
    if "可选产品" in header or "入云" in text:
        return "optional_product"
    return "business_attribute"


def is_service_attribute_row(item):
    if is_blank_business_row(item) or is_supporting_form_business_row(item):
        return False
    if item.get("category") == "说明":
        return False
    if is_base_package_row(item) or row_matches(item, INCLUDED_ITEM_KEYWORDS):
        return False
    if is_optional_service_row(item) or is_fee_rule_row(item):
        return False
    text = row_text(item)
    return bool(text and (item.get("description") or checkbox_options(text)))


def normalize_service_attribute(item):
    description = item.get("description", "")
    return {
        "attribute_type": service_attribute_type(item),
        "name": item.get("name") or item.get("category") or "",
        "category": item.get("category", ""),
        "value": description,
        "options": checkbox_options(description),
        "source_evidence": row_text(item),
        "row_index": item.get("row_index"),
    }


def extract_service_attributes(main_rows):
    attributes = []
    for item in classify_business_rows(main_rows):
        if is_service_attribute_row(item):
            attributes.append(normalize_service_attribute(item))
    return dedupe(attributes)


def build_document(path):
    doc = Document(str(path))
    paras = paragraphs(doc)
    tables = doc.tables
    main_rows = get_rows(tables[0]) if tables else []
    all_rows = flatten_table_rows(tables)
    business_rows = classify_business_rows(all_rows or main_rows)
    all_business_text = " ".join(item["description"] for item in business_rows)

    included_items = [
        item for item in business_rows
        if row_matches(item, INCLUDED_ITEM_KEYWORDS)
    ]
    optional_services = [
        item for item in business_rows
        if is_optional_service_row(item)
    ]
    fees = [
        item for item in business_rows
        if is_fee_rule_row(item)
    ]

    base_package_rows = [item for item in business_rows if is_base_package_row(item)]
    base_packages = parse_base_package_rows(base_package_rows)
    if not base_packages:
        base_package_text = " ".join(item["description"] for item in base_package_rows)
        base_packages = parse_base_packages(base_package_text or all_business_text)
    base_packages.extend(parse_variable_package_rows(business_rows))
    uplink_packages = parse_uplink_packages(all_business_text)
    if uplink_packages and not any(item["name"] == "上行升速包" for item in optional_services):
        optional_services.append({
            "row_index": None,
            "category": "可选增值服务",
            "name": "上行升速包",
            "description": "、".join(f"{p['name']} {p['amount']}元/{p['billing_period']}" for p in uplink_packages),
            "options": [f"{p['name']} {p['amount']}元/{p['billing_period']}" for p in uplink_packages],
        })

    meta = extract_meta(path, paras)
    warnings = []
    if not any(field["value"] for field in extract_required_fields(main_rows)):
        warnings.append("文档是空白申请登记表，未发现已填写的客户值；当前结果主要反映套餐规则和可选项。")
    if not base_packages:
        warnings.append("未解析到明确的基础套餐价格项，建议人工复核原表格。")

    extracted = {
        "document_type": "carrier_product_application_form",
        **meta,
        "customer_application_fields": extract_required_fields(main_rows),
        "filled_customer_values": [],
        "base_packages": base_packages,
        "included_items": dedupe(included_items),
        "service_attributes": extract_service_attributes(main_rows),
        "optional_services": dedupe(optional_services),
        "fees": dedupe(fees),
        "sla": extract_sla(tables, business_rows),
        "supporting_forms": extract_supporting_forms(tables),
        "marketing_rules": extract_marketing_rules(paras),
        "raw_business_rows": business_rows,
        "quality": {
            "overall_score": 0.82 if base_packages else 0.68,
            "method": "docx_table_rule_extraction",
            "warnings": warnings,
            "suggestions": [
                "如果后续配置 LLM API Key，可用项目原 LangGraph 链路做语义字段补全和冲突校验。",
                "业务验收时建议重点核对基础套餐价格、协议期、一次性费用和可选增值服务。"
            ],
        },
        "summary": "",
    }
    extracted["summary"] = (
        f"{meta['title']}，识别到 {len(base_packages)} 个基础套餐价格项、"
        f"{len(optional_services)} 个可选/增值服务项、{len(fees)} 条费用相关规则。"
    )
    return extracted


def iter_docx_files(data_dir=DATA_DIR):
    return [
        path
        for path in sorted(data_dir.glob("*.docx"))
        if not path.name.startswith("~$")
    ]


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    documents = []
    for index, path in enumerate(iter_docx_files(), 1):
        document = build_document(path)
        document["id"] = f"doc-{index}"
        documents.append(document)
        safe_name = f"document_{index}.json"
        (OUTPUT_DIR / safe_name).write_text(
            json.dumps(document, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

    bundle = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
            "source_directory": str(DATA_DIR.relative_to(ROOT)),
        "document_count": len(documents),
        "notes": [
            "本结果由 agent.product_doc_agent 本地规则抽取生成。",
            "当前环境未配置 LLM API Key，因此结果来自 DOCX 表格规则抽取，适合先给业务人员做可用性预览。"
        ],
        "documents": documents,
    }
    (OUTPUT_DIR / "all_documents.json").write_text(
        json.dumps(bundle, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps({
        "document_count": len(documents),
        "outputs": [str((OUTPUT_DIR / f"document_{i}.json").relative_to(ROOT)) for i in range(1, len(documents) + 1)],
        "bundle": str((OUTPUT_DIR / "all_documents.json").relative_to(ROOT)),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
